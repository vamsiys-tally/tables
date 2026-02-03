"""
Report Generator Module.

Generates Excel error reports for evaluation results.

Report Structure:
    Sheet 1: Summary - File-level metrics
    Sheet 2: Unmatched GT - Ground truth rows not extracted
    Sheet 3: Mismatched - Extracted rows with errors
    Sheet 4: Side-by-Side - Field-by-field comparison

Example:
    >>> generator = ReportGenerator()
    >>> generator.generate_file_report(comparison_result, "output/errors.xlsx")
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional, Any
import logging

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from tests.evaluation.comparator import (
    ComparisonResult,
    TransactionMatch,
    MismatchType,
    FieldComparison,
)
from tests.evaluation.metrics import AccuracyMetrics

logger = logging.getLogger(__name__)


# =============================================================================
# Styles
# =============================================================================

if HAS_OPENPYXL:
    # Header style
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")

    # Error highlight
    ERROR_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ERROR_FONT = Font(color="9C0006")

    # Success highlight
    SUCCESS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    SUCCESS_FONT = Font(color="006100")

    # Warning highlight
    WARNING_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    WARNING_FONT = Font(color="9C5700")

    # Border
    THIN_BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )


# =============================================================================
# Report Generator
# =============================================================================

class ReportGenerator:
    """
    Generates Excel reports for evaluation results.

    Creates detailed error reports with multiple sheets for analysis.

    Example:
        >>> generator = ReportGenerator()
        >>> generator.generate_file_report(result, "errors/hdfc_001_errors.xlsx")
    """

    def __init__(self, output_dir: Optional[str | Path] = None):
        """
        Initialize the report generator.

        Args:
            output_dir: Default output directory for reports
        """
        if not HAS_OPENPYXL:
            raise ImportError(
                "openpyxl is required for report generation. "
                "Install with: pip install openpyxl"
            )

        self.output_dir = Path(output_dir) if output_dir else Path(".")

    def generate_file_report(
        self,
        result: ComparisonResult,
        output_path: Optional[str | Path] = None,
    ) -> Path:
        """
        Generate an error report for a single file comparison.

        Args:
            result: ComparisonResult from comparator
            output_path: Output path for Excel file

        Returns:
            Path to generated report

        Creates 4 sheets:
            1. Summary - Overall metrics
            2. Unmatched GT - Ground truth rows not extracted
            3. Mismatched - Extracted rows with errors
            4. Side-by-Side - Field-by-field comparison
        """
        if output_path is None:
            file_stem = Path(result.file_name).stem
            output_path = self.output_dir / f"{file_stem}_errors.xlsx"
        else:
            output_path = Path(output_path)

        # Ensure output directory exists
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()

        # Sheet 1: Summary
        self._create_summary_sheet(wb.active, result)
        wb.active.title = "Summary"

        # Sheet 2: Unmatched GT rows
        ws_unmatched = wb.create_sheet("Unmatched GT")
        self._create_unmatched_gt_sheet(ws_unmatched, result)

        # Sheet 3: Mismatched extracted rows
        ws_mismatched = wb.create_sheet("Mismatched")
        self._create_mismatched_sheet(ws_mismatched, result)

        # Sheet 4: Side-by-side comparison
        ws_comparison = wb.create_sheet("Side-by-Side")
        self._create_comparison_sheet(ws_comparison, result)

        wb.save(output_path)
        logger.info(f"Generated report: {output_path}")

        return output_path

    def _create_summary_sheet(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
        result: ComparisonResult,
    ) -> None:
        """Create the summary sheet."""
        # Title
        ws["A1"] = "Evaluation Summary"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:B1")

        # Metrics
        metrics = [
            ("File", result.file_name),
            ("", ""),
            ("Ground Truth Transactions", result.gt_count),
            ("Extracted Transactions", result.extracted_count),
            ("Matched Transactions", result.matched_count),
            ("Missed (in GT, not extracted)", result.missed_count),
            ("Extra (extracted, not in GT)", result.extra_count),
            ("", ""),
            ("Accuracy", f"{result.accuracy:.2%}"),
            ("Perfect Match", "Yes" if result.is_perfect else "No"),
        ]

        for row_idx, (label, value) in enumerate(metrics, start=3):
            ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            cell = ws.cell(row=row_idx, column=2, value=value)

            # Highlight accuracy
            if label == "Accuracy":
                if result.accuracy >= 1.0:
                    cell.fill = SUCCESS_FILL
                elif result.accuracy >= 0.9:
                    cell.fill = WARNING_FILL
                else:
                    cell.fill = ERROR_FILL

        # Error breakdown
        error_breakdown = result.get_error_breakdown()
        if error_breakdown:
            start_row = len(metrics) + 5
            ws.cell(row=start_row, column=1, value="Error Breakdown").font = Font(
                bold=True, size=12
            )

            for idx, (error_type, count) in enumerate(
                sorted(error_breakdown.items(), key=lambda x: -x[1])
            ):
                ws.cell(row=start_row + idx + 1, column=1, value=error_type)
                ws.cell(row=start_row + idx + 1, column=2, value=count)

        # Adjust column widths
        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 50

    def _create_unmatched_gt_sheet(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
        result: ComparisonResult,
    ) -> None:
        """Create sheet for unmatched ground truth rows."""
        headers = [
            "GT_Row#",
            "Date",
            "Value Date",
            "Description",
            "Reference",
            "Debit",
            "Credit",
            "Balance",
            "Issue",
        ]

        # Write headers
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT

        # Write unmatched rows
        unmatched = result.get_missed()
        row_idx = 2

        for match in unmatched:
            if match.gt_transaction is None:
                continue

            gt = match.gt_transaction
            ws.cell(row=row_idx, column=1, value=gt.row_number)
            ws.cell(
                row=row_idx,
                column=2,
                value=gt.date.isoformat() if gt.date else "",
            )
            ws.cell(
                row=row_idx,
                column=3,
                value=gt.value_date.isoformat() if gt.value_date else "",
            )
            ws.cell(row=row_idx, column=4, value=gt.description)
            ws.cell(row=row_idx, column=5, value=gt.reference or "")
            ws.cell(row=row_idx, column=6, value=str(gt.debit) if gt.debit else "")
            ws.cell(row=row_idx, column=7, value=str(gt.credit) if gt.credit else "")
            ws.cell(row=row_idx, column=8, value=str(gt.balance) if gt.balance else "")
            ws.cell(row=row_idx, column=9, value=match.mismatch_type.value)

            # Highlight row
            for col in range(1, 10):
                ws.cell(row=row_idx, column=col).fill = ERROR_FILL

            row_idx += 1

        # Adjust column widths
        self._auto_adjust_columns(ws)

        if row_idx == 2:
            ws.cell(row=2, column=1, value="No unmatched ground truth rows")

    def _create_mismatched_sheet(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
        result: ComparisonResult,
    ) -> None:
        """Create sheet for mismatched/extra extracted rows."""
        headers = [
            "Ext_Row#",
            "Date",
            "Description",
            "Debit",
            "Credit",
            "Balance",
            "Issue",
            "Matched_GT_Row",
        ]

        # Write headers
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT

        # Get mismatched rows (not perfect matches and not pure GT misses)
        mismatched = [
            m for m in result.matches
            if not m.is_match and m.extracted_transaction is not None
        ]

        row_idx = 2
        for match in mismatched:
            ext = match.extracted_transaction

            ws.cell(row=row_idx, column=1, value=match.extracted_row)
            ws.cell(
                row=row_idx,
                column=2,
                value=str(ext.get("transaction_date", ext.get("date", ""))),
            )
            ws.cell(row=row_idx, column=3, value=ext.get("description", ""))
            ws.cell(row=row_idx, column=4, value=str(ext.get("debit", "")))
            ws.cell(row=row_idx, column=5, value=str(ext.get("credit", "")))
            ws.cell(row=row_idx, column=6, value=str(ext.get("balance", "")))
            ws.cell(row=row_idx, column=7, value=match.mismatch_type.value)
            ws.cell(
                row=row_idx,
                column=8,
                value=match.gt_row if match.gt_row else "-",
            )

            # Highlight based on issue type
            fill = ERROR_FILL if match.mismatch_type == MismatchType.EXTRA_ROW else WARNING_FILL
            for col in range(1, 9):
                ws.cell(row=row_idx, column=col).fill = fill

            row_idx += 1

        self._auto_adjust_columns(ws)

        if row_idx == 2:
            ws.cell(row=2, column=1, value="No mismatched extracted rows")

    def _create_comparison_sheet(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
        result: ComparisonResult,
    ) -> None:
        """Create side-by-side comparison sheet."""
        headers = ["GT_Row", "Field", "Ground Truth", "Extracted", "Match"]

        # Write headers
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT

        # Get matches with field comparisons (excluding perfect matches and pure misses)
        with_comparisons = [
            m for m in result.matches
            if not m.is_match
            and m.field_comparisons
            and m.gt_transaction is not None
            and m.extracted_transaction is not None
        ]

        row_idx = 2
        for match in with_comparisons:
            gt_row = match.gt_row

            for fc in match.field_comparisons:
                ws.cell(row=row_idx, column=1, value=gt_row)
                ws.cell(row=row_idx, column=2, value=fc.field_name)
                ws.cell(
                    row=row_idx,
                    column=3,
                    value=str(fc.gt_value) if fc.gt_value is not None else "",
                )
                ws.cell(
                    row=row_idx,
                    column=4,
                    value=str(fc.extracted_value) if fc.extracted_value is not None else "",
                )
                ws.cell(row=row_idx, column=5, value="✓" if fc.matches else "✗")

                # Highlight mismatched fields
                if not fc.matches:
                    for col in range(1, 6):
                        ws.cell(row=row_idx, column=col).fill = ERROR_FILL

                row_idx += 1

            # Add blank row between transactions
            row_idx += 1

        self._auto_adjust_columns(ws)

        if row_idx == 2:
            ws.cell(row=2, column=1, value="No field comparisons available")

    def _auto_adjust_columns(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
    ) -> None:
        """Auto-adjust column widths based on content."""
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass

            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

    def generate_summary_report(
        self,
        metrics: AccuracyMetrics,
        output_path: str | Path,
    ) -> Path:
        """
        Generate a summary report for all evaluated files.

        Args:
            metrics: AccuracyMetrics with aggregate results
            output_path: Output path for Excel file

        Returns:
            Path to generated report
        """
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()

        # Sheet 1: Overall Summary
        ws_summary = wb.active
        ws_summary.title = "Summary"
        self._create_overall_summary(ws_summary, metrics)

        # Sheet 2: File Results
        ws_files = wb.create_sheet("File Results")
        self._create_file_results_sheet(ws_files, metrics)

        # Sheet 3: By Bank (if available)
        if metrics.by_bank:
            ws_bank = wb.create_sheet("By Bank")
            self._create_category_sheet(ws_bank, metrics.by_bank, "Bank")

        # Sheet 4: Error Analysis
        ws_errors = wb.create_sheet("Error Analysis")
        self._create_error_analysis_sheet(ws_errors, metrics)

        wb.save(output_path)
        logger.info(f"Generated summary report: {output_path}")

        return output_path

    def _create_overall_summary(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
        metrics: AccuracyMetrics,
    ) -> None:
        """Create overall summary sheet."""
        ws["A1"] = "Evaluation Run Summary"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:B1")

        data = [
            ("", ""),
            ("Files", ""),
            ("Total Files", metrics.total_files),
            ("Perfect Files (100%)", metrics.perfect_files),
            ("Files with Errors", metrics.total_files - metrics.perfect_files),
            ("", ""),
            ("Transactions", ""),
            ("Ground Truth", metrics.total_gt_transactions),
            ("Extracted", metrics.total_extracted_transactions),
            ("Matched", metrics.total_matched_transactions),
            ("Missed", metrics.total_missed_transactions),
            ("Extra", metrics.total_extra_transactions),
            ("", ""),
            ("Overall Accuracy", f"{metrics.overall_accuracy:.2%}"),
        ]

        if metrics.duration_seconds:
            data.append(("Duration", f"{metrics.duration_seconds:.1f}s"))

        for row_idx, (label, value) in enumerate(data, start=3):
            cell_label = ws.cell(row=row_idx, column=1, value=label)
            cell_value = ws.cell(row=row_idx, column=2, value=value)

            if label in ("Files", "Transactions"):
                cell_label.font = Font(bold=True, size=11)
            elif label:
                cell_label.font = Font(bold=True)

            if label == "Overall Accuracy":
                if metrics.overall_accuracy >= 1.0:
                    cell_value.fill = SUCCESS_FILL
                elif metrics.overall_accuracy >= 0.95:
                    cell_value.fill = WARNING_FILL
                else:
                    cell_value.fill = ERROR_FILL

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 20

    def _create_file_results_sheet(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
        metrics: AccuracyMetrics,
    ) -> None:
        """Create file results sheet."""
        headers = [
            "File",
            "GT Count",
            "Extracted",
            "Matched",
            "Missed",
            "Extra",
            "Accuracy",
            "Status",
        ]

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

        # Sort by accuracy (worst first)
        sorted_files = sorted(metrics.file_metrics, key=lambda f: f.accuracy)

        for row_idx, file_metric in enumerate(sorted_files, start=2):
            ws.cell(row=row_idx, column=1, value=file_metric.file_name)
            ws.cell(row=row_idx, column=2, value=file_metric.gt_count)
            ws.cell(row=row_idx, column=3, value=file_metric.extracted_count)
            ws.cell(row=row_idx, column=4, value=file_metric.matched_count)
            ws.cell(row=row_idx, column=5, value=file_metric.missed_count)
            ws.cell(row=row_idx, column=6, value=file_metric.extra_count)
            ws.cell(row=row_idx, column=7, value=f"{file_metric.accuracy:.2%}")
            ws.cell(
                row=row_idx,
                column=8,
                value="PERFECT" if file_metric.is_perfect else "ERRORS",
            )

            # Highlight based on accuracy
            if file_metric.is_perfect:
                fill = SUCCESS_FILL
            elif file_metric.accuracy >= 0.9:
                fill = WARNING_FILL
            else:
                fill = ERROR_FILL

            for col in range(1, 9):
                ws.cell(row=row_idx, column=col).fill = fill

        self._auto_adjust_columns(ws)

    def _create_category_sheet(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
        category_metrics: dict[str, Any],
        category_name: str,
    ) -> None:
        """Create category breakdown sheet."""
        headers = [
            category_name,
            "Files",
            "Perfect",
            "GT Total",
            "Matched",
            "Missed",
            "Extra",
            "Accuracy",
        ]

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

        for row_idx, (name, cat) in enumerate(
            sorted(category_metrics.items()), start=2
        ):
            ws.cell(row=row_idx, column=1, value=name)
            ws.cell(row=row_idx, column=2, value=cat.file_count)
            ws.cell(row=row_idx, column=3, value=cat.perfect_count)
            ws.cell(row=row_idx, column=4, value=cat.total_gt)
            ws.cell(row=row_idx, column=5, value=cat.total_matched)
            ws.cell(row=row_idx, column=6, value=cat.total_missed)
            ws.cell(row=row_idx, column=7, value=cat.total_extra)
            ws.cell(row=row_idx, column=8, value=f"{cat.accuracy:.2%}")

        self._auto_adjust_columns(ws)

    def _create_error_analysis_sheet(
        self,
        ws: "openpyxl.worksheet.worksheet.Worksheet",
        metrics: AccuracyMetrics,
    ) -> None:
        """Create error analysis sheet."""
        headers = ["Error Type", "Count", "Percentage"]

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

        total_errors = sum(metrics.error_breakdown.values())

        for row_idx, (error_type, count) in enumerate(
            sorted(metrics.error_breakdown.items(), key=lambda x: -x[1]), start=2
        ):
            ws.cell(row=row_idx, column=1, value=error_type)
            ws.cell(row=row_idx, column=2, value=count)
            percentage = (count / total_errors * 100) if total_errors > 0 else 0
            ws.cell(row=row_idx, column=3, value=f"{percentage:.1f}%")

        self._auto_adjust_columns(ws)
